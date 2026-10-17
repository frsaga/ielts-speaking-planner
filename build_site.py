from __future__ import annotations
import argparse,json
from pathlib import Path
ROOT=Path(__file__).resolve().parent
TEMPLATE=ROOT/'.index.template.html'
def text(value):
    return '' if value is None else str(value).strip()


def rows_from_sheet(sheet_name: str, part: str, workbook):
    ws = workbook[sheet_name]
    questions = []
    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = [text(item) for item in row[:6]]
        if not any(values):
            continue
        if part == 'P1':
            source, topic, number, english, chinese, original_status = values
            prompts = ''
            group = topic
            chinese_title = ''
        elif part == 'P2':
            source, number, chinese_title, english, prompts, original_status = values
            topic = chinese_title
            group = chinese_title
            chinese = chinese_title
        else:
            source, group, number, english, chinese, original_status = values
            topic = group.split('｜')[-1] if '｜' in group else group
            prompts = ''
            chinese_title = ''
        if not english:
            continue
        questions.append({
            'key': f'{part}-{row_index}',
            'part': part,
            'number': number,
            'topic': topic,
            'group': group,
            'source': source,
            'question': english,
            'chinese': chinese,
            'chineseTitle': chinese_title,
            'prompts': prompts,
            'initialStatus': original_status or '未练',
        })
    return questions


def main():
    parser=argparse.ArgumentParser(description="Build a speaking practice planner from local data.")
    parser.add_argument('--questions-json',type=Path,default=ROOT/'demo-questions.json')
    parser.add_argument('--workbook',type=Path)
    args=parser.parse_args()
    if args.workbook:
        from openpyxl import load_workbook
        workbook=load_workbook(args.workbook,read_only=True,data_only=True)
        questions=[]
        for part in ('P1','P2','P3'):
            sheet=next((name for name in workbook.sheetnames if name.startswith(part)),None)
            if sheet: questions.extend(rows_from_sheet(sheet,part,workbook))
    else:
        questions=json.loads(args.questions_json.read_text())
    if not questions or len({q['key'] for q in questions})!=len(questions):
        raise ValueError('Provide at least one question with unique keys.')
    payload=json.dumps(questions,ensure_ascii=False).replace('</','<\\/')
    html=TEMPLATE.read_text().replace('__QUESTION_DATA__',payload)
    (ROOT/'index.html').write_text(html)
    print(f'Built {len(questions)} public/sample questions.')
if __name__=='__main__': main()
